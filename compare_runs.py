#!/usr/bin/env python3
"""
Ski Run Comparison Tool
Creates an Excel file comparing two ski runs with color coding.

Usage:
    python compare_runs.py

Output:
    Creates 'ski_comparison.xlsx' with full analysis
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

# =============================================================================
# DATA - Edit these values with your times
# =============================================================================

SKIER1_NAME = "Wissting"
SKIER2_NAME = "Jakobsen"

# Split times for each skier (in seconds)
# First value is START (will be marked as such)
SKIER1_SPLITS = [
    "START", 1.32, 1.00, 1.32, 1.24, 1.32, 1.00, 1.24, 1.16, 1.36,
    0.32, 1.20, 1.16, 1.16, 1.16, 1.08, 0.36, 0.36, 1.20, 1.12,
    1.12, 1.00, 1.12, 1.12, 1.16, 1.04, 1.12, 1.00, 1.12
]

SKIER2_SPLITS = [
    "START", 1.28, 1.00, 1.28, 1.20, 1.32, 1.32, 1.28, 1.16, 1.24,
    0.28, 1.12, 1.08, 1.16, 1.16, 1.04, 0.32, 0.36, 1.16, 1.12,
    1.08, 1.08, 1.00, 1.12, 1.16, 1.08, 1.16, 1.04, 1.12
]

# =============================================================================
# SCRIPT - Don't edit below unless you know what you're doing
# =============================================================================

def create_comparison_excel():
    """Create a complete comparison Excel file."""

    wb = Workbook()
    ws = wb.active
    ws.title = "Run Comparison"

    # Styles
    header_font = Font(bold=True, size=12, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    center = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

    # Headers
    headers = ["Gate", SKIER1_NAME, SKIER2_NAME, "Diff (s)", "Diff (%)", "Faster"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin_border

    # Data rows
    num_gates = min(len(SKIER1_SPLITS), len(SKIER2_SPLITS))

    total_s1 = 0
    total_s2 = 0
    s1_faster = 0
    s2_faster = 0
    equal = 0

    for i in range(num_gates):
        row = i + 2
        gate = i + 1
        s1 = SKIER1_SPLITS[i]
        s2 = SKIER2_SPLITS[i]

        # Gate number
        ws.cell(row=row, column=1, value=gate).alignment = center
        ws.cell(row=row, column=1).border = thin_border

        # Skier 1 split
        if s1 == "START":
            ws.cell(row=row, column=2, value="START").alignment = center
        else:
            ws.cell(row=row, column=2, value=s1).alignment = center
            ws.cell(row=row, column=2).number_format = '0.00'
            total_s1 += s1
        ws.cell(row=row, column=2).border = thin_border

        # Skier 2 split
        if s2 == "START":
            ws.cell(row=row, column=3, value="START").alignment = center
        else:
            ws.cell(row=row, column=3, value=s2).alignment = center
            ws.cell(row=row, column=3).number_format = '0.00'
            total_s2 += s2
        ws.cell(row=row, column=3).border = thin_border

        # Calculate difference
        if s1 == "START" or s2 == "START":
            ws.cell(row=row, column=4, value="-").alignment = center
            ws.cell(row=row, column=5, value="-").alignment = center
            ws.cell(row=row, column=6, value="-").alignment = center
        else:
            diff = round(s1 - s2, 2)
            diff_pct = round((s1 - s2) / s2 * 100, 1)

            # Diff column
            diff_cell = ws.cell(row=row, column=4, value=diff)
            diff_cell.alignment = center
            diff_cell.number_format = '0.00'
            diff_cell.border = thin_border

            # Color coding for diff
            if diff > 0:
                diff_cell.fill = red_fill  # Skier 1 slower
                s2_faster += 1
            elif diff < 0:
                diff_cell.fill = green_fill  # Skier 1 faster
                s1_faster += 1
            else:
                diff_cell.fill = yellow_fill  # Equal
                equal += 1

            # Diff % column
            pct_cell = ws.cell(row=row, column=5, value=f"{diff_pct}%")
            pct_cell.alignment = center
            pct_cell.border = thin_border
            if diff > 0:
                pct_cell.fill = red_fill
            elif diff < 0:
                pct_cell.fill = green_fill
            else:
                pct_cell.fill = yellow_fill

            # Faster column
            if diff > 0:
                faster = SKIER2_NAME
            elif diff < 0:
                faster = SKIER1_NAME
            else:
                faster = "Lika"
            faster_cell = ws.cell(row=row, column=6, value=faster)
            faster_cell.alignment = center
            faster_cell.border = thin_border

        ws.cell(row=row, column=4).border = thin_border
        ws.cell(row=row, column=5).border = thin_border
        ws.cell(row=row, column=6).border = thin_border

    # Total row
    total_row = num_gates + 2
    total_diff = round(total_s1 - total_s2, 2)
    total_diff_pct = round((total_s1 - total_s2) / total_s2 * 100, 1) if total_s2 > 0 else 0

    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
    ws.cell(row=total_row, column=1).alignment = center
    ws.cell(row=total_row, column=2, value=round(total_s1, 2)).font = Font(bold=True)
    ws.cell(row=total_row, column=2).alignment = center
    ws.cell(row=total_row, column=2).number_format = '0.00'
    ws.cell(row=total_row, column=3, value=round(total_s2, 2)).font = Font(bold=True)
    ws.cell(row=total_row, column=3).alignment = center
    ws.cell(row=total_row, column=3).number_format = '0.00'

    total_diff_cell = ws.cell(row=total_row, column=4, value=total_diff)
    total_diff_cell.font = Font(bold=True)
    total_diff_cell.alignment = center
    total_diff_cell.number_format = '0.00'
    if total_diff > 0:
        total_diff_cell.fill = red_fill
    elif total_diff < 0:
        total_diff_cell.fill = green_fill

    ws.cell(row=total_row, column=5, value=f"{total_diff_pct}%").font = Font(bold=True)
    ws.cell(row=total_row, column=5).alignment = center

    winner = SKIER2_NAME if total_diff > 0 else (SKIER1_NAME if total_diff < 0 else "Lika")
    ws.cell(row=total_row, column=6, value=winner).font = Font(bold=True)
    ws.cell(row=total_row, column=6).alignment = center

    # Add border to total row
    for col in range(1, 7):
        ws.cell(row=total_row, column=col).border = thin_border

    # Summary section
    summary_row = total_row + 3
    ws.cell(row=summary_row, column=1, value="SAMMANFATTNING").font = Font(bold=True, size=14)

    ws.cell(row=summary_row + 2, column=1, value=f"{SKIER1_NAME} snabbast:")
    ws.cell(row=summary_row + 2, column=2, value=f"{s1_faster} portar")

    ws.cell(row=summary_row + 3, column=1, value=f"{SKIER2_NAME} snabbast:")
    ws.cell(row=summary_row + 3, column=2, value=f"{s2_faster} portar")

    ws.cell(row=summary_row + 4, column=1, value="Lika:")
    ws.cell(row=summary_row + 4, column=2, value=f"{equal} portar")

    ws.cell(row=summary_row + 6, column=1, value="Total skillnad:")
    ws.cell(row=summary_row + 6, column=2, value=f"{abs(total_diff):.2f}s")

    ws.cell(row=summary_row + 7, column=1, value="Vinnare:")
    winner_cell = ws.cell(row=summary_row + 7, column=2, value=winner)
    winner_cell.font = Font(bold=True, size=12)

    # Column widths
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 10
    ws.column_dimensions['E'].width = 10
    ws.column_dimensions['F'].width = 12

    # Save
    filename = "ski_comparison.xlsx"
    wb.save(filename)
    print(f"\n✅ Fil skapad: {filename}")
    print(f"\n📊 RESULTAT:")
    print(f"   {SKIER1_NAME}: {total_s1:.2f}s")
    print(f"   {SKIER2_NAME}: {total_s2:.2f}s")
    print(f"   Skillnad: {abs(total_diff):.2f}s")
    print(f"   Vinnare: {winner}")
    print(f"\n   {SKIER1_NAME} snabbast: {s1_faster} portar")
    print(f"   {SKIER2_NAME} snabbast: {s2_faster} portar")
    print(f"   Lika: {equal} portar")

    return filename

if __name__ == "__main__":
    print("🎿 Ski Run Comparison Tool")
    print("=" * 40)
    create_comparison_excel()
    print("\n" + "=" * 40)
    print("Öppna filen i Excel eller Google Sheets!")
